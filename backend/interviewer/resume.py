"""Resume ingestion: uploaded file -> plain text -> structured ResumeProfile (Pydantic)."""
import io
import re

MAX_BYTES = 5 * 1024 * 1024   # 5 MB
MAX_CHARS = 14_000            # text passed to the model
MIN_CHARS = 80


class ResumeError(ValueError):
    """A problem the user can fix (wrong type, scanned image, too big). Shown to them as-is."""


def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if len(data) > MAX_BYTES:
        raise ResumeError("That file is larger than 5 MB.")

    if name.endswith(".pdf"):
        try:
            from pypdf import PdfReader
        except ImportError as err:
            raise ResumeError("PDF support isn't installed on the server. Run: pip install pypdf") from err
        try:
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages[:8])
        except Exception as err:  # noqa: BLE001
            raise ResumeError("That PDF couldn't be read. It may be damaged or password protected.") from err
    elif name.endswith(".docx"):
        try:
            import docx
        except ImportError as err:
            raise ResumeError("Word support isn't installed on the server. Run: pip install python-docx") from err
        try:
            doc = docx.Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        except Exception as err:  # noqa: BLE001
            raise ResumeError("That Word file couldn't be read.") from err
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ResumeError("Please upload a PDF, DOCX or TXT file.")

    text = re.sub(r"[ \t]+", " ", text.replace("\x00", ""))
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < MIN_CHARS:
        raise ResumeError(
            "We couldn't find readable text in that file. If it's a scan or an image, upload a text-based PDF or DOCX."
        )
    return text[:MAX_CHARS]


def parse_resume(llm, text: str):
    """Structured extraction with the same self-correcting retry loop used elsewhere."""
    from . import prompts
    from .schemas import ResumeProfile

    runnable = llm.with_structured_output(ResumeProfile)
    messages = [("system", prompts.RESUME_SYSTEM), ("human", prompts.build_resume_prompt(text))]
    last_err = None
    for _ in range(3):
        try:
            result = runnable.invoke(messages)
            return result if isinstance(result, ResumeProfile) else ResumeProfile.model_validate(result)
        except Exception as err:  # noqa: BLE001
            last_err = err
            messages = messages + [("human", f"Your previous output was invalid: {err}\nReturn a corrected response.")]
    raise ResumeError("We couldn't understand that resume. Please try again or use a different file.") from last_err